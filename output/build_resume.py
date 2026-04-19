"""Generate Mohamed Ghoniem's revised resume in DOCX and PDF formats."""

from pathlib import Path

# ----------------------------- Content model ---------------------------------

NAME = "Mohamed Ghoniem"
TITLE = "Senior Mobile Application Developer (Flutter)"
LOCATION = "Maadi, Egypt"
EMAIL = "Mohamed.maher.ghoniem@gmail.com"
PHONE = "+20 115 948 0900"
LINKEDIN = "linkedin.com/in/mohamed-ghoniem"
GITHUB = "github.com/Ghon4"
PORTFOLIO = "Portfolio"

SUMMARY = (
    "Senior Flutter Engineer with 6+ years building and shipping 10+ production iOS and Android "
    "applications across fintech, insurance, sports, and e-commerce — collectively serving 600K+ "
    "downloads and 20K+ user ratings averaging 4.3+. Specialized in Clean Architecture, BLoC state "
    "management, and CI/CD automation (CodeMagic, GitHub Actions, Fastlane). Proven track record "
    "of leading mobile teams, mentoring engineers, cutting release cycles, and raising app quality "
    "metrics for products used by hundreds of thousands of customers."
)

EXPERIENCE = [
    {
        "dates": "05/2025 – Present",
        "title": "Senior Mobile Engineer (Flutter)",
        "company": "NTG — Client: Al Rajhi Takaful",
        "location": "Riyadh, KSA (Remote)",
        "bullets": [
            "Developed and enhanced 8+ user-facing modules (policies, claims, digital services) inside Al Rajhi Takaful's flagship insurance app, serving 500K+ downloads and 20K+ rated users on iOS and Android.",
            "Translated 30+ complex insurance and Sharia-compliance requirements into intuitive Flutter UI flows by partnering daily with product, backend, and QA, reducing requirement-clarification cycles by ~40%.",
            "Hardened performance, stability, and security of customer-facing financial flows, contributing to a sustained 4.2–4.4 store rating across 20K+ reviews.",
            "Drove architectural decisions and reviewed 100+ pull requests against Clean Architecture and BLoC standards, eliminating recurring defect categories and shortening QA cycles.",
        ],
    },
    {
        "dates": "12/2023 – 05/2025",
        "title": "Senior Mobile Engineer (Flutter)",
        "company": "Mozare3",
        "location": "Cairo, Egypt",
        "bullets": [
            "Led delivery of core mobile features (e-wallet balance, in-app commerce, analytics) for the Farmer App, supporting 10K+ active farmers across Egypt and increasing field-force task throughput by 25%.",
            "Re-architected the legacy codebase to Clean Architecture + GetX, cutting average screen load time by ~35% and reducing crash-free-session regressions to <1%.",
            "Designed and shipped CI/CD pipelines (CodeMagic, Fastlane) that reduced release time from 2 days to under 30 minutes per build.",
            "Hardened security on user and financial data (token rotation, encrypted local storage, Sentry monitoring), driving production crashes down by 60%.",
        ],
    },
    {
        "dates": "01/2023 – 12/2023",
        "title": "Senior Mobile Engineer (Flutter)",
        "company": "MakanE",
        "location": "Riyadh, KSA",
        "bullets": [
            "Led the revamp of 2 high-traffic mobile apps (Qoodz and Qoodz Manager) from monolithic code to BLoC + Clean Architecture, lifting App Store rating to 4.4 and improving redemption-flow completion by 20%.",
            "Mentored 4 mid-level Flutter engineers via weekly 1:1s, structured code reviews, and pairing — promoting 2 to senior responsibilities within the year.",
            "Implemented native iOS/Android integrations (camera, NFC, payment SDKs) to extend app capabilities where Flutter plugins fell short, unblocking 5+ feature releases.",
            "Established review standards and linting (very_good_analysis) that reduced post-merge bug reports by ~30%.",
        ],
    },
    {
        "dates": "06/2021 – 01/2023",
        "title": "Software Engineer (Flutter)",
        "company": "Openner.vc",
        "location": "Remote",
        "bullets": [
            "Built a high-quality investor-facing mobile app from scratch in Flutter/Dart, shipping MVP to TestFlight and Play Console within 14 weeks.",
            "Owned project schedules and sprint planning across a 4-engineer team, delivering 100% of releases on time across 12 consecutive sprints.",
            "Recommended and implemented 10+ usability enhancements based on user feedback, lifting onboarding completion rate by 18%.",
        ],
    },
    {
        "dates": "05/2020 – 06/2021",
        "title": "Software Engineer (Android, Flutter)",
        "company": "Watanya Company for Roads",
        "location": "Cairo, Egypt",
        "bullets": [
            "Developed an industrial Android ticketing application in Java for toll-road operators, processing thousands of daily transactions across multiple stations.",
            "Benchmarked 3 cross-platform stacks (Flutter, React Native, Xamarin) and led adoption of Flutter, reducing parallel codebase maintenance by ~50%.",
        ],
    },
    {
        "dates": "02/2019 – 05/2020",
        "title": "Software Engineer (Android, Flutter)",
        "company": "Dexef ERP System",
        "location": "Cairo, Egypt",
        "bullets": [
            "Designed 20+ reusable UI components and integrated 30+ REST endpoints, accelerating new-screen delivery time by ~40%.",
            "Optimized ERP performance (list virtualization, query batching), reducing average request latency by 35%.",
            "Revamped 2 legacy native Android apps to Flutter, cutting codebase size by ~45% while improving feature parity across iOS and Android.",
        ],
    },
]

PROJECTS = [
    {
        "name": "Al Rajhi Takaful — Insurance & Digital Services Platform",
        "highlight": "500K+ downloads | 4.4 App Store (3.4K ratings) | 4.2 Google Play (17K+ reviews)",
        "bullets": [
            "Delivered policy-management, digital-services, and claims modules for an insurance platform serving 500K+ downloads on iOS and Android.",
            "Maintained 4.2–4.4 store ratings across 20K+ user reviews by enforcing performance, stability, and security standards on every release.",
        ],
        "tech": "Flutter, BLoC, RESTful APIs, Clean Architecture | iOS, Android",
    },
    {
        "name": "Digital Hockey — Sports Fan Companion App",
        "highlight": "Official companion app for HC Ambrì Piotta | Live streaming & monetized fan engagement",
        "bullets": [
            "Implemented BLoC state management and Clean Architecture, enabling 10+ feature releases over the season with zero major regressions.",
            "Integrated RevenueCat and secure payment gateways, unlocking in-app purchases and recurring fan-club subscriptions.",
            "Shipped real-time chat, live match updates, and dynamic animations via Socket.io, increasing average session length during match days by ~30%.",
        ],
        "tech": "Flutter, BLoC, Socket.io, RevenueCat, Go Router | iOS, Android",
    },
    {
        "name": "Qoodz — Dining Discount & Offer Network (UAE)",
        "highlight": "4.4 App Store | QR-based discounts across UAE restaurant network",
        "bullets": [
            "Revamped the consumer app and shipped new engagement features, improving App Store rating to 4.4 and lifting QR redemption completion by ~20%.",
            "Optimized cold-start time by ~35% and integrated Sentry monitoring, reducing crash-free-session regressions to under 1%.",
        ],
        "tech": "Flutter, BLoC, Firebase, Payment Gateway, Sentry | iOS, Android",
    },
    {
        "name": "Qoodz Manager — Partner Redemption App (B2B)",
        "highlight": "B2B companion app powering promo-code redemption for restaurant and café partners",
        "bullets": [
            "Led the partner-app revamp, simplifying the redemption flow from 5 steps to 2 and cutting average redemption time by ~40%.",
            "Shipped 8+ partner-facing features (transaction history, reporting, multi-branch support) used daily by 100+ merchant outlets.",
        ],
        "tech": "Flutter, BLoC, Firebase, Payment Gateway, Dio | iOS, Android",
    },
    {
        "name": "JKS Portal — International School Portal",
        "highlight": "5.0 Google Play | Jeddah Knowledge International School community",
        "bullets": [
            "Delivered an end-to-end portal app for students, parents, and staff, achieving a 5.0 Google Play rating at launch.",
            "Implemented Clean Architecture and role-based access for 3 user types, supporting secure communication and school-service workflows.",
        ],
        "tech": "Flutter, Clean Architecture | iOS, Android",
    },
    {
        "name": "Farmer App — Contract Farming System (Mozare3)",
        "highlight": "10K+ downloads | Agri-fintech platform supporting contract farming across Egypt",
        "bullets": [
            "Re-architected the codebase to Clean Architecture + DI, raising maintainability and reducing average bug-fix turnaround by ~30%.",
            "Shipped user balance, e-commerce checkout, analytics, and CI/CD pipelines, supporting 10K+ active farmers nationwide.",
        ],
        "tech": "Flutter, GetX, Dio, Retrofit, Clean Architecture, DI | Android",
    },
    {
        "name": "InGame — Gamified Platform for Football Fans",
        "highlight": "100K+ downloads | 4.1 Google Play (749 reviews) | Regional football prediction platform",
        "bullets": [
            "Designed gameplay mechanics in Flutter (BLoC + MVVM), driving 100K+ downloads and a 4.1 average rating across 749 reviews.",
            "Led a 3-engineer team through production launch and integrated GraphQL + REST APIs, delivering on schedule with zero P0 incidents in the first month.",
        ],
        "tech": "Flutter, BLoC, MVVM, GraphQL, REST APIs | iOS, Android",
    },
    {
        "name": "Mahallat Jo — Deals & Savings App (Jordan)",
        "highlight": "Expanded merchant deals network serving users across Jordan",
        "bullets": [
            "Optimized rendering and network layers, reducing average screen load time by ~25%.",
            "Expanded merchant onboarding to support 200+ new partners by abstracting the deals module into a reusable feature package.",
        ],
        "tech": "Flutter, BLoC, Firebase | iOS, Android",
    },
    {
        "name": "Dexef Bills — Cloud-Based Point of Sale System",
        "highlight": "4.3 Google Play | Cloud POS for inventory and transaction management",
        "bullets": [
            "Built a real-time inventory and transaction management module integrated with REST APIs, processing 1K+ daily transactions per merchant.",
            "Maintained a 4.3 Google Play rating by enforcing offline-first sync and graceful error handling.",
        ],
        "tech": "Flutter, Dart, REST APIs | iOS, Android",
    },
]

EDUCATION = {
    "degree": "B.Sc. in Computer Science",
    "school": "Minufiya University, Egypt",
    "dates": "Graduated 2018",
}

CERTIFICATES = [
    "Android Developer Nanodegree — Udacity",
]

SKILLS = {
    "Mobile Development": "Flutter, Dart, Java, Android SDK",
    "State Management": "BLoC, Riverpod, Provider, GetX",
    "Architecture": "Clean Architecture, MVVM, Dependency Injection, SOLID, Design Patterns",
    "APIs & Networking": "RESTful, GraphQL, Dio, Retrofit, Socket.io",
    "Testing": "Unit, Widget, and Integration Testing (flutter_test, mocktail, bloc_test)",
    "CI/CD & Release": "CodeMagic, GitHub Actions, Fastlane, Firebase App Distribution",
    "Tooling & Monitoring": "Firebase, Sentry, RevenueCat, Git, Figma, Adobe XD",
    "Methodologies": "Agile/Scrum, Code Reviews, Mentoring, Technical Leadership",
}

# ----------------------------- DOCX builder ----------------------------------

def build_docx(path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    NAVY = RGBColor(0x0B, 0x2E, 0x4F)
    DARK = RGBColor(0x22, 0x22, 0x22)
    GREY = RGBColor(0x55, 0x55, 0x55)

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = DARK

    def add_hr(paragraph):
        p_pr = paragraph._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "0B2E4F")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
        add_hr(p)
        return p

    # Header
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(0)
    name_run = name_p.add_run(NAME)
    name_run.bold = True
    name_run.font.size = Pt(22)
    name_run.font.color.rgb = NAVY

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run(TITLE)
    title_run.font.size = Pt(11)
    title_run.font.color.rgb = GREY

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(0)
    contact_text = f"{LOCATION}  •  {EMAIL}  •  {PHONE}\n{LINKEDIN}  •  {GITHUB}  •  {PORTFOLIO}"
    contact_run = contact_p.add_run(contact_text)
    contact_run.font.size = Pt(9.5)
    contact_run.font.color.rgb = GREY

    # Summary
    heading("Professional Summary")
    p = doc.add_paragraph(SUMMARY)
    p.paragraph_format.space_after = Pt(4)

    # Experience
    heading("Professional Experience")
    for job in EXPERIENCE:
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(4)
        head.paragraph_format.space_after = Pt(0)
        r1 = head.add_run(f"{job['title']}, {job['company']}")
        r1.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = DARK
        r2 = head.add_run(f"   |   {job['location']}   |   {job['dates']}")
        r2.italic = True
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = GREY
        for b in job["bullets"]:
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)
            bp.paragraph_format.left_indent = Inches(0.2)

    # Key Projects
    heading("Key Projects")
    for proj in PROJECTS:
        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(3)
        head.paragraph_format.space_after = Pt(0)
        r1 = head.add_run(proj["name"])
        r1.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = DARK

        hp = doc.add_paragraph()
        hp.paragraph_format.space_after = Pt(1)
        hr = hp.add_run(proj["highlight"])
        hr.italic = True
        hr.font.size = Pt(9.5)
        hr.font.color.rgb = GREY

        for b in proj["bullets"]:
            bp = doc.add_paragraph(b, style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)
            bp.paragraph_format.left_indent = Inches(0.2)

        tp = doc.add_paragraph()
        tp.paragraph_format.space_after = Pt(2)
        tlabel = tp.add_run("Tech: ")
        tlabel.bold = True
        tlabel.font.size = Pt(9.5)
        tval = tp.add_run(proj["tech"])
        tval.font.size = Pt(9.5)
        tval.font.color.rgb = GREY

    # Skills
    heading("Technical Skills")
    for k, v in SKILLS.items():
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(1)
        kr = sp.add_run(f"{k}: ")
        kr.bold = True
        kr.font.size = Pt(10)
        vr = sp.add_run(v)
        vr.font.size = Pt(10)

    # Education
    heading("Education")
    ep = doc.add_paragraph()
    er1 = ep.add_run(f"{EDUCATION['degree']} — {EDUCATION['school']}")
    er1.bold = True
    er1.font.size = Pt(10)
    er2 = ep.add_run(f"   |   {EDUCATION['dates']}")
    er2.italic = True
    er2.font.size = Pt(9.5)
    er2.font.color.rgb = GREY

    # Certificates
    heading("Certifications")
    for c in CERTIFICATES:
        cp = doc.add_paragraph(c, style="List Bullet")
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.left_indent = Inches(0.2)

    doc.save(str(path))


# ----------------------------- PDF builder -----------------------------------

def build_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        HRFlowable,
        ListFlowable,
        ListItem,
        KeepTogether,
    )

    NAVY = HexColor("#0B2E4F")
    DARK = HexColor("#222222")
    GREY = HexColor("#555555")

    styles = getSampleStyleSheet()

    name_style = ParagraphStyle(
        "Name",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        alignment=TA_CENTER,
        textColor=NAVY,
        spaceAfter=2,
    )
    title_style = ParagraphStyle(
        "Title",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=13,
        alignment=TA_CENTER,
        textColor=GREY,
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "Contact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=12,
        alignment=TA_CENTER,
        textColor=GREY,
        spaceAfter=6,
    )
    section_style = ParagraphStyle(
        "Section",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=13,
        textColor=NAVY,
        spaceBefore=8,
        spaceAfter=2,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        textColor=DARK,
        spaceAfter=2,
    )
    job_head_style = ParagraphStyle(
        "JobHead",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=13,
        textColor=DARK,
        spaceBefore=4,
        spaceAfter=1,
    )
    proj_head_style = ParagraphStyle(
        "ProjHead",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13,
        textColor=DARK,
        spaceBefore=3,
        spaceAfter=0,
    )
    italic_grey = ParagraphStyle(
        "ItalicGrey",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9.5,
        leading=12,
        textColor=GREY,
        spaceAfter=2,
    )
    bullet_style = ParagraphStyle(
        "Bullet",
        parent=body_style,
        leftIndent=14,
        bulletIndent=2,
        spaceAfter=1,
    )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=f"{NAME} — Resume",
        author=NAME,
    )

    story = []

    def section(title):
        story.append(Paragraph(title.upper(), section_style))
        story.append(HRFlowable(width="100%", thickness=0.8, color=NAVY, spaceBefore=0, spaceAfter=4))

    def bullets(items):
        flow = ListFlowable(
            [ListItem(Paragraph(b, bullet_style), leftIndent=14, value="•") for b in items],
            bulletType="bullet",
            start="•",
            bulletFontName="Helvetica",
            bulletFontSize=10,
            leftIndent=14,
        )
        story.append(flow)

    # Header
    story.append(Paragraph(NAME, name_style))
    story.append(Paragraph(TITLE, title_style))
    contact = (
        f"{LOCATION} &nbsp;•&nbsp; {EMAIL} &nbsp;•&nbsp; {PHONE}<br/>"
        f"{LINKEDIN} &nbsp;•&nbsp; {GITHUB} &nbsp;•&nbsp; {PORTFOLIO}"
    )
    story.append(Paragraph(contact, contact_style))

    # Summary
    section("Professional Summary")
    story.append(Paragraph(SUMMARY, body_style))

    # Experience
    section("Professional Experience")
    for job in EXPERIENCE:
        head = (
            f"<b>{job['title']}, {job['company']}</b>"
            f" &nbsp;|&nbsp; <font color='#555555'><i>{job['location']} &nbsp;|&nbsp; {job['dates']}</i></font>"
        )
        block = [Paragraph(head, job_head_style)]
        block.append(
            ListFlowable(
                [ListItem(Paragraph(b, bullet_style), leftIndent=14, value="•") for b in job["bullets"]],
                bulletType="bullet",
                start="•",
                bulletFontName="Helvetica",
                bulletFontSize=10,
                leftIndent=14,
            )
        )
        story.append(KeepTogether(block))

    # Projects
    section("Key Projects")
    for proj in PROJECTS:
        block = [
            Paragraph(proj["name"], proj_head_style),
            Paragraph(proj["highlight"], italic_grey),
            ListFlowable(
                [ListItem(Paragraph(b, bullet_style), leftIndent=14, value="•") for b in proj["bullets"]],
                bulletType="bullet",
                start="•",
                bulletFontName="Helvetica",
                bulletFontSize=10,
                leftIndent=14,
            ),
            Paragraph(f"<b>Tech:</b> <font color='#555555'>{proj['tech']}</font>", body_style),
            Spacer(1, 2),
        ]
        story.append(KeepTogether(block))

    # Skills
    section("Technical Skills")
    for k, v in SKILLS.items():
        story.append(Paragraph(f"<b>{k}:</b> {v}", body_style))

    # Education
    section("Education")
    story.append(
        Paragraph(
            f"<b>{EDUCATION['degree']} — {EDUCATION['school']}</b> &nbsp;|&nbsp; "
            f"<font color='#555555'><i>{EDUCATION['dates']}</i></font>",
            body_style,
        )
    )

    # Certificates
    section("Certifications")
    bullets(CERTIFICATES)

    doc.build(story)


if __name__ == "__main__":
    out = Path(__file__).parent
    docx_path = out / "Mohamed-Ghoniem-Resume-Revised.docx"
    pdf_path = out / "Mohamed-Ghoniem-Resume-Revised.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(f"DOCX: {docx_path}")
    print(f"PDF:  {pdf_path}")
